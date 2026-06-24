"""
Generic report exporters following the SKILLS.md xlsx/pdf toolchains:
- xlsx via openpyxl (Arial, bold header fill, zeros formatted as "-")
- pdf via reportlab (Helvetica, A4)

Both consume the Report dict shape produced by reporting.services and return raw
bytes for an HTTP download.
"""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
)
from reportlab.lib.styles import getSampleStyleSheet

HEADER_FILL = "0E1726"  # Mizan ink (design token)
# Excel number format: 2dp, thousands separators, zeros shown as "-".
NUM_FORMAT = "#,##0.00;-#,##0.00;\"-\""


def _fmt_num(value):
    if value is None:
        return "-"
    return f"{Decimal(value):,.2f}"


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def to_xlsx(report):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = report["title"][:31]

    arial = "Arial"
    columns = report["columns"]
    ncols = len(columns)

    # Title + subtitle.
    ws.cell(row=1, column=1, value=report["title"]).font = Font(name=arial, bold=True, size=14)
    ws.cell(row=2, column=1, value=report["subtitle"]).font = Font(name=arial, size=9, color="677084")

    # Column headers.
    header_row = 4
    for c, col in enumerate(columns, 1):
        cell = ws.cell(row=header_row, column=c, value=col["label"])
        cell.font = Font(name=arial, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="right" if col["numeric"] else "left")

    # Data rows.
    r = header_row + 1
    for row in report["rows"]:
        style = row.get("style", "normal")
        bold = style in ("total", "subtotal", "section")
        for c, (col, value) in enumerate(zip(columns, row["cells"]), 1):
            cell = ws.cell(row=r, column=c)
            if col["numeric"]:
                if value is None:
                    cell.value = None
                else:
                    cell.value = float(Decimal(value))
                    cell.number_format = NUM_FORMAT
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.value = value
            cell.font = Font(name=arial, bold=bold)
        r += 1

    # Column widths.
    for c, col in enumerate(columns, 1):
        width = 38 if not col["numeric"] else 16
        ws.column_dimensions[get_column_letter(c)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def to_pdf(report):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
                            leftMargin=16 * mm, rightMargin=16 * mm)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(report["title"], styles["Title"]),
        Paragraph(report["subtitle"], styles["Normal"]),
        Spacer(1, 8 * mm),
    ]

    columns = report["columns"]
    table_data = [[col["label"] for col in columns]]
    for row in report["rows"]:
        cells = []
        for col, value in zip(columns, row["cells"]):
            cells.append(_fmt_num(value) if col["numeric"] else (value or ""))
        table_data.append(cells)

    table = Table(table_data, repeatRows=1, hAlign="LEFT")
    ts = [
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0E1726")),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 9),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E4E7EC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
    ]
    # Right-align numeric columns; bold total/subtotal/section rows.
    for ci, col in enumerate(columns):
        if col["numeric"]:
            ts.append(("ALIGN", (ci, 1), (ci, -1), "RIGHT"))
    for ri, row in enumerate(report["rows"], start=1):
        if row.get("style") in ("total", "subtotal", "section"):
            ts.append(("FONT", (0, ri), (-1, ri), "Helvetica-Bold", 9))
    table.setStyle(TableStyle(ts))

    elements.append(table)
    doc.build(elements)
    return buf.getvalue()


EXTENSIONS = {
    "xlsx": ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", to_xlsx),
    "pdf": ("pdf", "application/pdf", to_pdf),
}
